#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-FileCopyrightText: 2026 BLogic Mikroelektronik - Berk Muammer Kuzu and Berkin Demircan
# SPDX-License-Identifier: GPL-3.0-only
# Licensed under the GNU General Public License version 3 only.
# See the LICENSE file in the repository root for the full license text.

# ============================================
# Ostim BLogic Mikroelektronik
# write_dtr_222.py  -  DTR 2.2.2 cevre birim bolumunu yazar
# ============================================
"""2.2.2 Çevre Birim Tasarım Detayları (8p) -> BLogic_MCU_DTR_taslak_01.docx"""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'BLogic_MCU_DTR_taslak_01.docx'
doc = docx.Document(SRC)

def _font(run, size, bold=False, italic=False):
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold; run.italic = italic

def make_para(text, size=11, bold=False, italic=False, sb=0, sa=6, justify=True):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(text), size, bold, italic)
    return p

def make_heading(text, sz=12, sb=10):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(3); pf.keep_with_next = True
    _font(p.add_run(text), sz, bold=True)
    return p

def make_img(label, desc):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(6); pf.space_after = Pt(0)
    _font(p.add_run('[ ' + label + ' ]'), 11, bold=True); place(p)
    q = doc.add_paragraph(style='Normal')
    qf = q.paragraph_format; qf.space_before = Pt(0); qf.space_after = Pt(6)
    _font(q.add_run(desc), 10, italic=True)
    return q

def set_cell(cell, text, bold=False, size=10, shade=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    _font(p.add_run(text), size, bold)
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), shade); tcPr.append(shd)

def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        b.append(e)
    tblPr.append(b)

def _width(table, pct=5000):
    tblPr = table._tbl.tblPr
    for ex in tblPr.findall(qn('w:tblW')): tblPr.remove(ex)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), str(pct)); tblPr.append(w)

def make_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _borders(t); _width(t); t.autofit = True
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, bold=True, size=10, shade='D9D9D9')
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(t.rows[i].cells[j], val, size=10)
    return t

# regmap kısayolu
def regmap(base, rows):
    return make_table(["Ofset", "Ad", "R/W", "Açıklama"], rows)

heading_p = None; placeholder = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 3' and 'Çevre Birim Tasarım Detayları' in p.text:
        heading_p = p
    if p.text.strip().startswith('{Şartnamede tanımlanan çevre birimlerin'):
        placeholder = p
assert heading_p is not None and placeholder is not None, 'bulunamadı'

anchor = heading_p._element
def place(obj):
    global anchor
    anchor.addnext(obj._element); anchor = obj._element; return obj

place(make_para(
    "BLogic MCU, şartname EK-2'de tanımlanan çevre birimi takımını barındırır: genel amaçlı "
    "UART_0, özel veri-akışı birimi UART_1, GPIO, Timer, I2C Master ve QSPI Master. Yedi çevre "
    "birimi de (YZ hızlandırıcı CSR dâhil) ortak bir AXI4-Lite slave şablonu üzerine "
    "tasarlanmış olup periph_decoder'ın arkasına bağlanır. Tüm birimlerin yazmaç düzenleri "
    "EK-2 ile uyumlu olacak biçimde belirlenmiş, C sürücü başlık dosyaları hazırlanmış ve her "
    "biri kendi öz-denetimli testiyle doğrulanmıştır. Şartnamenin bu aşamada zorunlu kıldığı "
    "asgari gereksinim olan UART, sisteme tümüyle entegre edilmiş ve DDK demo-testbench'inde "
    "doğrulanmıştır. Aşağıda her çevre birimi tasarım yaklaşımı, yazmaç haritası, doğrulama "
    "durumu ve kısıtlarıyla ayrıntılandırılmaktadır."))

place(make_heading("Ortak Tasarım Yaklaşımı"))
place(make_para(
    "Bütün çevre birimleri aynı el-yapımı AXI4-Lite slave FSM şablonunu kullanır: 32-bit "
    "adres/veri, 4-bit yazma maskesi (wstrb), tek-bekleyen erişim ve OKAY (2'b00) yanıt. Bu "
    "tek-tip yaklaşım, doğrulamayı (her arayüze bağlanan SVA protokol monitörü) ve yeni birim "
    "eklemeyi kolaylaştırır; yazmaç erişimleri öngörülebilir ve birbiriyle tutarlıdır. Verilog "
    "kaynak dosyaları rtl/peripherals/ altında, C sürücü başlıkları ise sw/drivers/ altında "
    "yer alır."))

# UART_0
place(make_heading("UART_0 — Genel Amaçlı UART  (taban 0x4000_0000)"))
place(make_para(
    "UART_0, kanıtlanmış açık kaynaklı Alex Forencich uart_tx/uart_rx çekirdeklerini saran, "
    "yapılandırılabilir baud hızına sahip bir seri haberleşme birimidir. Baud hızı CPB yazmacıyla "
    "belirlenir (baud = clk / CPB; 50 MHz için varsayılan CPB=434 → 115200). Donanıma giden "
    "ön-bölme CPB>>3 olarak uygulanır. Dur biti sayısı STP ile (1 / 1.5 / 2) seçilebilir ve "
    "20-bit bir uzatma sayacıyla gerçeklenir. Birimin kesmesi yoktur; durum CFG yazmacındaki "
    "bayraklarla yoklanır. FIFO bulunmaz; tek-bayt TDR/RDR yazmaçları üzerinden çalışır. UART_0 "
    "tümüyle entegre olup teknotest 'hello world' senaryosunu ve baud-tarama (baud-sweep) "
    "testini geçmiştir."))
place(regmap("0x4000_0000", [
    ["0x00", "CPB", "RW", "Bağds bölücü (baud = clk/CPB; reset 434 = 50 MHz/115200)"],
    ["0x04", "STP", "RW", "Dur biti [1:0]: 00=1, 01=1.5, 1x=2"],
    ["0x08", "RDR", "RO", "Alınan bayt"],
    ["0x0C", "TDR", "RW", "Gönderilecek bayt"],
    ["0x10", "CFG", "RW", "[0] tx_en, [1] rx_done, [2] tx_done"],
]))
make_img("GÖRSEL - önerilen: Çevre birimi doğrulama kanıtı",
    "UART TX/RX simülasyon dalga formu veya teknotest 'hello world' konsol çıktısı; istenirse "
    "I2C işlemi (START/ADR/ACK/DATA/STOP) ve QSPI flash okuma dalga formları da eklenebilir.")

# UART_1
place(make_heading("UART_1 — Veri-Akışı (Stream / DMA) Birimi  (taban 0x4000_0300)"))
place(make_para(
    "UART_1, UART_0 ile aynı seri çekirdeği barındırır; ek olarak RX hattından gelen baytları "
    "CPU'yu meşgul etmeden doğrudan AI SRAM'e taşıyan özgün bir DMA mekanizmasına sahiptir. "
    "Baytlar little-endian biçimde 32-bit sözcüklere paketlenir ve birimin kendi AXI4 yaz-only "
    "master'ı (ID 3) ile 4-bayt'lık beat'ler hâlinde AI SRAM'e (varsayılan 0x0003_0000) "
    "yazılır. Aktarım tamamlandığında bir çevrimlik kesme darbesi üretilir (bit 18). Akış "
    "etkinken arbiter sahipliği bu birime verilir. ABORT, devam eden AXI yazması bittikten "
    "sonra sahipliği güvenli biçimde bırakır. Temel UART yazmaçlarına (0x00-0x10) ek olarak:"))
place(regmap("0x4000_0300", [
    ["0x14", "STRM_ADDR", "RW", "DMA hedef adresi (varsayılan 0x0003_0000 = AI SRAM girişi)"],
    ["0x18", "STRM_LEN", "RW", "Aktarılacak bayt sayısı (varsayılan 1960 = 49×40 INT8)"],
    ["0x1C", "STRM_CTRL", "WO", "[0] START, [1] ABORT (darbe)"],
    ["0x20", "STRM_STAT", "RO", "[0] BUSY, [1] DONE, [31:16] alınan bayt sayısı"],
]))

# GPIO
place(make_heading("GPIO  (taban 0x4000_0100)"))
place(make_para(
    "GPIO birimi 16 sabit giriş ve 16 sabit çıkış hattı sağlar; yön sabit olduğundan ayrı bir "
    "yön yazmacı yoktur. Girişler iki kademeli flip-flop ile saat alanına senkronize edilir "
    "(metastabilite koruması). FPGA'da çıkışlar LED ve Pmod'a, girişler anahtar ve butonlara "
    "bağlanır. Birimin kesmesi yoktur. GPIO LED testiyle doğrulanmıştır."))
place(regmap("0x4000_0100", [
    ["0x00", "IDR", "RO", "Giriş verisi [15:0] (senkronize)"],
    ["0x04", "ODR", "RW", "Çıkış verisi [15:0]"],
]))

# Timer
place(make_heading("Timer  (taban 0x4000_0200)"))
place(make_para(
    "Timer, prescaler'lı 32-bit bir sayaçtır; sayaç her (PRE+1) çevrimde bir artar/azalır. "
    "ARE değeri otomatik yeniden-yükleme/karşılaştırma sınırıdır; sayaç bu değere ulaşınca "
    "sıfırlanır ve olay sayacı (EVN) artar. MOD biti yukarı/aşağı sayım yönünü belirler. Kesme "
    "seviye-tetiklemelidir: timer_irq_o, EVN sıfırdan farklı olduğu sürece aktiftir (bit 16) ve "
    "yazılım EVC'ye yazarak temizler. Timer sistem testinde doğrulanmıştır."))
place(regmap("0x4000_0200", [
    ["0x00", "PRE", "RW", "Prescaler (her PRE+1 çevrimde tik)"],
    ["0x04", "ARE", "RW", "Auto-reload / karşılaştırma değeri (reset 0xFFFFFFFF)"],
    ["0x08", "CLR", "RW", "[0]=1 sayacı sıfırla"],
    ["0x0C", "ENA", "RW", "[0]=1 etkinleştir"],
    ["0x10", "MOD", "RW", "[0]: 1=yukarı, 0=aşağı sayım"],
    ["0x14", "CNT", "RO", "Anlık sayaç değeri"],
    ["0x18", "EVN", "RO", "Olay (taşma) sayacı"],
    ["0x1C", "EVC", "RW", "[0]=1 olay sayacını ve kesmeyi temizle"],
]))

# I2C
place(make_heading("I2C Master  (taban 0x4000_0400)"))
place(make_para(
    "I2C Master, 400 kHz hızında çalışan bit-tabanlı (4 fazlı) bir denetleyicidir (QDIV=31). "
    "SCL push-pull, SDA ise açık-drain sürülür (sda_oe; 1 = SDA'yı aşağı çek). FSM "
    "IDLE/START/BITS/ACK/STOP aşamalarından oluşur; START koşulu SCL yüksekken SDA'nın "
    "düşmesiyle, STOP ise yükselmesiyle üretilir. Bir işlemde 1-4 bayt aktarılabilir (NBY); 7-bit "
    "slave adresi ADR ile verilir. NACK durumunda STOP üretilip done ve nack_err bayrakları "
    "kurulur. I2C sistem testiyle doğrulanmıştır. Mevcut sürümde clock stretching ve repeated "
    "START desteklenmez (bkz. eksik/planlanan)."))
place(regmap("0x4000_0400", [
    ["0x00", "NBY", "RW", "Bayt sayısı (1-4; 0→1, >4→4 olarak kırpılır)"],
    ["0x04", "ADR", "RW", "7-bit slave adresi [6:0]"],
    ["0x08", "RDR", "RO", "Alınan veri (ilk bayt [7:0])"],
    ["0x0C", "TDR", "RW", "Gönderilecek veri (ilk bayt [7:0])"],
    ["0x10", "CFG", "RW", "[0] tx_en, [1] tx_done, [2] rx_en, [3] rx_done, [4] NACK"],
]))

# QSPI
place(make_heading("QSPI Master  (taban 0x4000_0500)"))
place(make_para(
    "QSPI Master, kart üzerindeki S25FL256S flash belleğiyle haberleşen ve x1/x2/x4 (tek/çift/"
    "dörtlü hat) veri modlarını destekleyen bir denetleyicidir; komut, adres ve dummy fazları "
    "her zaman x1'dir. Adres fazı 3 veya 4 baytlık olabilir (FCR[2]) ve adres fazı davranışı "
    "FCR[4:3] ile (AUTO/FORCE/SKIP) denetlenir. CCR yazmacına yazıldığında işlem başlar; veri "
    "64-derin TX/RX FIFO'ları üzerinden DR ile taşınır. Boot sırasında bootloader bu birimi "
    "kullanarak flash'tan READ (0x03) komutuyla program imajını Buyruk SRAM'e kopyalar. QSPI, "
    "çalışma modları testi ve boot ile doğrulanmıştır."))
place(regmap("0x4000_0500", [
    ["0x00", "CCR", "RW", "Komut yapılandırma: instr[7:0], data_mode[9:8] (x1/x2/x4), dir[10], dummy[15:11], data_len[23:16], prescaler[30:25]; yazınca işlem başlar"],
    ["0x04", "ADR", "RW", "Flash adresi"],
    ["0x08", "DR", "RW", "Veri (64-derin TX/RX FIFO)"],
    ["0x0C", "STA", "RO", "[0] done, [1] busy, FIFO boş/dolu bayrakları"],
    ["0x10", "FCR", "RW", "[0] RX flush, [1] TX flush, [2] 3/4-bayt adres, [4:3] adres-faz modu"],
]))

# C sürücüleri
place(make_heading("C Sürücüleri (Driver Başlıkları)"))
place(make_para(
    "Tüm çevre birimleri için C sürücü başlık dosyaları hazırlanmıştır. sw/drivers/blogic_mcu.h "
    "dosyası yedi çevre biriminin taban adreslerini ve bellek-haritalı yazmaç yapılarını "
    "(UART_TypeDef, GPIO, Timer, I2C, QSPI, YZ-CSR) tanımlar; sürücü taban adresleri ve yazmaç "
    "ofsetleri RTL çözücüyle birebir tutarlıdır. sw/drivers/qspi.h ise QSPI için CCR oluşturma "
    "ve okuma yardımcılarını içerir. Bu başlıklar, sw/tests/ altındaki öz-denetimli testlerde "
    "(uart_hello, uart_baud_sweep, gpio_led_test, i2c_system_test, qspi_modes_test vb.) "
    "kullanılmaktadır."))

# durum özeti
place(make_heading("Tasarım ve Doğrulama Durumu"))
place(make_table(
    ["Çevre birimi", "Tasarım", "Doğrulama", "Durum / kısıt"],
    [
     ["UART_0", "Tamamlandı", "teknotest hello-world + baud-sweep + loopback", "Tam entegre (zorunlu); FIFO yok (tek-bayt)"],
     ["UART_1 (akış/DMA)", "Tamamlandı", "uart-stream öz-denetimli test", "DMA→AI SRAM çalışır; kesme (bit 18)"],
     ["GPIO", "Tamamlandı", "gpio_led_test", "16 giriş / 16 çıkış (sabit yön)"],
     ["Timer", "Tamamlandı", "sistem testi", "Seviye-tetiklemeli kesme (bit 16)"],
     ["I2C Master", "Tamamlandı", "i2c_system_test", "400 kHz; clock-stretch / repeated-START yok (planlı)"],
     ["QSPI Master", "Tamamlandı", "qspi_modes_test + boot", "x1/x2/x4; S25FL256S flash boot"],
    ]))

# artı/eksi
place(make_heading("Seçilen Yöntemlerin Artı ve Eksileri"))
place(make_table(
    ["Tasarım kararı", "Artıları", "Eksileri / risk"],
    [
     ["Ortak AXI4-Lite slave şablonu",
      "Tek-tip, doğrulaması ve genişletmesi kolay, öngörülebilir register erişimi",
      "Birim-bazlı optimize değil (ör. UART'ta FIFO yok)"],
     ["Forencich açık-kaynak UART çekirdekleri",
      "Sahada kanıtlanmış, geliştirme süresi kısa",
      "Tek-bayt; STP/CFG için sarmalayıcı gerekti"],
     ["Yoklama (poll) tabanlı UART/I2C/QSPI",
      "Basit, deterministik, kesme yükü yok",
      "CPU meşgul-bekler; yalnız Timer/YZ/akış kesme kullanır"],
     ["I2C bit-bang 4-fazlı motor",
      "Küçük alan, sade denetim",
      "Clock-stretch / repeated-START yok"],
     ["QSPI x1/x2/x4 + 64-derin FIFO",
      "Esnek, hızlı flash erişimi ve boot",
      "Daha karmaşık denetim mantığı"],
    ]))

# zorluklar
place(make_heading("Karşılaşılan Zorluklar ve Çözümleri"))
for t in [
    "1) UART zamanlama/baud uyumu: dur-biti süresinin tam tutturulması için STP uzatma sayacı "
    "eklendi; 50 MHz saat için CPB=434 ile 115200 bağds elde edildi.",
    "2) Gerçek S25FL256S flash'ta tCO kaynaklı bit kayması: QSPI RX örnekleme noktası yükselen "
    "kenara taşındı (SPI mod-0 düzeltmesi). Ayrıca FPGA STARTUPE2 ilk kenarları için bootloader'a "
    "warm-up okuması eklendi.",
    "3) FPGA'da I2C açık-drain SDA: gerçek açık-drain pad yerine tristate emülasyonu (sda_oe ? 0 "
    ": Z) ve pull-up ile çözüldü.",
    "4) UART_1 akış biriminde little-endian paketleme ve güvenli iptal: ABORT, devam eden AXI "
    "yazması tamamlandıktan sonra sahipliği bırakacak biçimde tasarlandı; kesme yalnız normal "
    "tamamlanmada üretilir.",
]:
    place(make_para(t, sa=4))

# eksik/planlanan
place(make_heading("Eksik / Planlanan Aktiviteler"))
place(make_para(
    "Mevcut sürümde tüm çevre birimleri tasarlanmış ve sisteme entegre edilmiştir; tamamlanmamış "
    "bir çevre birimi yoktur. Birkaç işlevsel iyileştirme bir sonraki aşamaya planlanmıştır: "
    "(i) I2C için clock stretching ve repeated START desteği eklenecektir (mevcut testler tekil "
    "yazma/okuma işlemleriyle geçmektedir); (ii) UART/I2C/QSPI için isteğe bağlı kesme hatları "
    "değerlendirilecektir (mevcut tasarımda bu birimler bilinçli olarak yoklama ile çalışır); "
    "(iii) UART'a opsiyonel TX/RX FIFO eklenmesi gözden geçirilecektir. Bu öğeler, asgari işlevin "
    "tamamlanmış olması gereksinimini etkilemez; UART tam olarak entegre ve doğrulanmış durumdadır."))

# placeholder'ı sil, kaydet
placeholder._element.getparent().remove(placeholder._element)
doc.save(SRC)
print("OK - 2.2.2 yazıldı, placeholder silindi.")
