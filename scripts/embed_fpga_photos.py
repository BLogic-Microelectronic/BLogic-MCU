#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-FileCopyrightText: 2026 BLogic Mikroelektronik - Berk Muammer Kuzu and Berkin Demircan
# SPDX-License-Identifier: GPL-3.0-only
# Licensed under the GNU General Public License version 3 only.
# See the LICENSE file in the repository root for the full license text.

"""2.4 FPGA bölümüne gerçek FPGA görsellerini gömer."""
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC = 'BLogic_MCU_DTR_taslak_01.docx'
doc = docx.Document(SRC)
CENTER = WD_ALIGN_PARAGRAPH.CENTER

def _font(run, size, bold=False, italic=False):
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold; run.italic = italic

def img_para(path, width_cm):
    p = doc.add_paragraph(style='Normal'); p.alignment = CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p

def cap_para(text):
    c = doc.add_paragraph(style='Normal'); c.alignment = CENTER
    c.paragraph_format.space_before = Pt(0); c.paragraph_format.space_after = Pt(8)
    _font(c.add_run(text), 9, italic=True)
    return c

def led_table(p_left, p_right, w=7.2):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, path in [(t.rows[0].cells[0], p_left), (t.rows[0].cells[1], p_right)]:
        cp = cell.paragraphs[0]; cp.alignment = CENTER
        cp.add_run().add_picture(path, width=Cm(w))
    return t

# --- anchors ---
periph_para = None; results_para = None
for p in doc.paragraphs:
    if p.text.strip().startswith('Çevre birimleri yazılımla donanım üzerinde sınanmıştır'):
        periph_para = p
    if p.text.strip().startswith('Tasarım, Genesys-2 kartında başarıyla sentezlenip'):
        results_para = p
assert periph_para is not None and results_para is not None, 'anchor bulunamadı'

# --- 2.4'teki [GÖRSEL - ZORUNLU] etiketleri + doc sonundaki ARTIK açıklama
#     paragraflarını (önceki make_img bug'ı) temizle ---
DESC_PREFIXES = (
    "İçermesi gerekenler:", "Yukarıdaki haritayı", "Reset -> Boot ROM",
    "sysclk_p/n (200 MHz LVDS)", "UART TX/RX simülasyon dalga formu",
    "report_utilization çıktısının", "Genesys-2 kartının fotoğrafı",
    "Host bilgisayardaki seri terminal",
)
removed = 0
for p in list(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('[ GÖRSEL - ZORUNLU') or t.startswith(DESC_PREFIXES):
        el = p._element; par = el.getparent()
        if par is not None:
            par.remove(el); removed += 1
print("silinen (etiket + artık açıklama):", removed)

class Cursor:
    def __init__(self, el): self.el = el
    def add(self, obj):
        self.el.addnext(obj._element); self.el = obj._element

# --- 1) Çevre birim testleri: UART + GPIO görselleri ---
c = Cursor(periph_para._element)
c.add(img_para('images/uart_hello.png', 15.0))
c.add(cap_para("Görsel 2.4.1 — FPGA'daki CV32E40P'den USB-UART üzerinden host terminale gelen "
               "'Hello World from BLogic MCU!' çıktısı (UART TX testi)."))
c.add(img_para('images/uart_calc.png', 12.0))
c.add(cap_para("Görsel 2.4.2 — UART üzerinden çalışan etkileşimli hesap makinesi: işlemci iki "
               "rakamı RX'ten okur, toplar ve sonucu TX'ten gönderir (çift yönlü UART testi)."))
c.add(led_table('images/gpio_ledon.jpg', 'images/gpio_ledoff.jpg', 7.0))
c.add(cap_para("Görsel 2.4.3 — GPIO çıkış testi: yazılımla sürülen LED'lerin Genesys-2 kartında "
               "açık (solda) ve kapalı (sağda) durumları."))

# --- 2) Alınan Sonuçlar: Vivado çıktıları ---
r = Cursor(results_para._element)
r.add(img_para('images/res_designruns.png', 16.0))
r.add(cap_para("Görsel 2.4.4 — Vivado Design Runs: sentez ve implementasyon tamamlandı, "
               "write_bitstream başarılı. WNS = +2.208 ns (zamanlama sağlandı), Failed Routes = 0; "
               "LUT 46.485 · FF 61.682 · BRAM 12.5 · DSP 17; toplam güç 0.487 W."))
r.add(img_para('images/res_util.png', 14.5))
r.add(cap_para("Görsel 2.4.5 — Post-implementation kaynak kullanımı: LUT %23, FF %15, BRAM %3, "
               "DSP %2, IO %9, BUFG %6, MMCM %10 (Kintex-7 XC7K325T)."))
r.add(img_para('images/res_floorplan.png', 7.5))
r.add(cap_para("Görsel 2.4.6 — Yerleştirme-bağlama sonrası implemented design görünümü; çip "
               "üzerinde kullanılan kaynaklar (mavi)."))

doc.save(SRC)
print("OK - 7 görsel gömüldü. inline_shapes:", len(doc.inline_shapes))
