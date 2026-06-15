#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2.4 FPGA Prototipleme (5p) -> BLogic_MCU_DTR_taslak_01.docx"""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'BLogic_MCU_DTR_taslak_01.docx'
doc = docx.Document(SRC)

def _font(run, size, bold=False, italic=False):
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold; run.italic = italic

def make_para(text, size=11, bold=False, italic=False, sb=0, sa=6, justify=True):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(text), size, bold, italic)
    return p

def make_heading(text):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(10); pf.space_after = Pt(3); pf.keep_with_next = True
    _font(p.add_run(text), 12, bold=True)
    return p

def make_img(label, desc):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(6); pf.space_after = Pt(0)
    _font(p.add_run('[ ' + label + ' ]'), 11, bold=True); place(p)
    q = doc.add_paragraph(style='Normal')
    qf = q.paragraph_format; qf.space_before = Pt(0); qf.space_after = Pt(6)
    _font(q.add_run(desc), 10, italic=True)
    return q

heading_p = None; placeholder = None
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') and 'FPGA Prototipleme' in p.text:
        heading_p = p
    if p.text.strip().startswith('{FPGA prototipleme detayları'):
        placeholder = p
assert heading_p is not None and placeholder is not None, 'bulunamadı'

anchor = heading_p._element
def place(obj):
    global anchor
    anchor.addnext(obj._element); anchor = obj._element; return obj

# ===================== İÇERİK =====================
place(make_para(
    "BLogic MCU, RTL tasarımının gerçek donanım üzerinde doğrulanması amacıyla Digilent "
    "Genesys-2 (Xilinx Kintex-7 XC7K325T-2FFG900C) kartında prototiplenmiştir. Kart-bağımsız "
    "soc_top RTL'i, FPGA'ya özgü saat/reset/IO uyarlamalarını içeren ince fpga_top "
    "sarmalayıcısıyla karta uyarlanmış; Vivado ile sentezlenip bitstream üretilmiş ve karta "
    "yüklenerek çalıştırılmıştır. İşlemci üzerinde bootloader ve uygulama yazılımı koşturulmuş, "
    "çevre birimleri donanım üzerinde sınanmıştır. Bu alt bölümde platform, üretim akışı, "
    "koşturulan yazılımlar, alınan sonuçlar ve karşılaşılan zorluklar ele alınmaktadır."))

place(make_heading("Donanım Platformu, Saat ve Reset"))
place(make_para(
    "Kart üzerindeki 200 MHz LVDS referans saat, IBUFDS ile tek-uçlu hale getirilip MMCME2_BASE "
    "ile 50 MHz sistem saatine dönüştürülür (×5 → VCO 1 GHz, ÷20 → 50 MHz) ve BUFG ile global ağ "
    "üzerinden dağıtılır. Reset, cpu_resetn butonu ile mmcm_locked sinyalinin mantıksal VE'siyle "
    "üretilip iki kademeli senkronizatörden geçirilir; böylece SoC, MMCM kilitlenene kadar "
    "reset'te tutulur (async-assert / sync-release). Kartın bring-up görünürlüğü için CPU ve "
    "reset'ten bağımsız serbest çalışan bir 'heartbeat' LED'i (~1.3 s periyot) eklenmiştir; bu "
    "LED'in yanıp sönmesi, saatin ve yapılandırmanın sağlıklı olduğunu anında gösterir. Ayrıntılı "
    "saat ve QSPI mimarisi 2.1'deki ilgili görselde verilmiştir."))

place(make_heading("Bitstream Üretimi ve Flash Programlama Akışı"))
place(make_para(
    "Bitstream, proje-dışı (non-project) bir Vivado betiği olan fpga/build_genesys2.tcl ile "
    "üretilir (hedef parça xc7k325tffg900-2). Betik, FPGA'da verimsiz sentezlenen latch-tabanlı "
    "kayıt dosyası yerine flip-flop tabanlı sürümü seçer, yalnız-doğrulamaya yönelik SVA "
    "monitörlerini sentezden hariç tutar ve yonga-üzeri bellekleri $readmemh ile Block RAM olarak "
    "çıkarır. İki açılış modu desteklenir: flash-boot (BOOT_ADDR=0x0 — Boot ROM içindeki "
    "bootloader program imajını QSPI flash'tan Buyruk SRAM'e kopyalar) ve SRAM-boot "
    "(BOOT_ADDR=0x10000 — firmware.hex doğrudan BRAM'e ön-yüklü olarak çalışır). Flash "
    "programlama fpga/flash_firmware.tcl ile yapılır: firmware .mcs biçimine dönüştürülüp kart "
    "üzerindeki S25FL256S flash'a yazılır ve doğrulanır."))

place(make_heading("İşlemci Üzerinde Yazılım Koşturma"))
place(make_para(
    "Evet — CV32E40P çekirdeği üzerinde gerçek yazılım koşturulmuştur. Flash-boot modunda, Boot "
    "ROM bootloader'ı QSPI üzerinden program imajını flash'tan Buyruk SRAM'e kopyalar ve "
    "uygulamaya atlar; SRAM-boot modunda uygulama BRAM'e ön-yüklü olarak doğrudan çalışır. Aynı C "
    "kaynakları (sw/ altında) hem Verilator simülasyonunda hem FPGA'da kullanılır ve rv32imc "
    "hedefiyle derlenir; böylece simülasyonda doğrulanan yazılım, değiştirilmeden donanımda "
    "koşturulur."))

place(make_heading("Çevre Birimlerinin Donanımda Testi"))
place(make_para(
    "Çevre birimleri yazılımla donanım üzerinde sınanmıştır. UART_0, kartın USB-UART (FT232) "
    "köprüsü üzerinden host bilgisayara bağlanmış ve uygulamanın gönderdiği 'Hello World' çıktısı "
    "seri terminalde gözlemlenmiştir. GPIO, anahtar/buton girişleri ile LED çıkışları üzerinden "
    "test edilmiş (anahtar durumunun LED'lere yansıması) ve heartbeat LED'i ile saat/reset "
    "bütünlüğü sürekli izlenmiştir. QSPI birimi, flash-boot akışının kartta çalışması sayesinde "
    "doğrudan doğrulanmıştır: program gerçek S25FL256S flash'tan okunup işlemci tarafından "
    "yürütülmektedir. Çevre birimlerine ait C sürücüleri sw/drivers/ altında yer almaktadır."))

place(make_heading("Alınan Sonuçlar"))
place(make_para(
    "Tasarım, Genesys-2 kartında başarıyla sentezlenip yapılandırılmış ve son derece kompakt bir "
    "kaynak profiliyle çalışmaktadır; Vivado kaynak kullanım (utilization) raporu LUT, flip-flop, "
    "BRAM, DSP ve MMCM kullanımını göstermektedir (aşağıdaki görsel). Kart üzerinde tasarımın "
    "çalıştığı ve UART üzerinden alınan çıktı, sonuçların görsellerle belgelenmesi gereği "
    "doğrultusunda aşağıda sunulmaktadır.", sa=4))
make_img("GÖRSEL - ZORUNLU: Vivado kaynak kullanım (utilization) raporu",
    "report_utilization çıktısının ekran görüntüsü: LUT / FF / BRAM36 / DSP48 / MMCM sayıları ve "
    "yüzdeleri. (Sayıları buradan al; özetteki 'Buraya değişecek' tahminlerini bu gerçek "
    "değerlerle güncelle.)")
make_img("GÖRSEL - ZORUNLU: Kart üzerinde çalışan tasarım",
    "Genesys-2 kartının fotoğrafı; tasarım koşarken (heartbeat LED'i yanarken / anahtar→LED "
    "yansıması görünür biçimde).")
make_img("GÖRSEL - ZORUNLU: FPGA'dan UART çıktısı",
    "Host bilgisayardaki seri terminal ekran görüntüsü: FPGA'daki işlemciden USB-UART üzerinden "
    "gelen 'Hello World from BLogic MCU!' çıktısı.")

place(make_heading("Karşılaşılan Zorluklar ve Çözümleri"))
for t in [
    "1) QSPI saati (SCLK) FPGA'da ayrılmış CCLK yapılandırma pininden geçtiği için normal bir IO "
    "olarak sürülemez. Çözüm: STARTUPE2/USRCCLKO ilkeli kullanıldı; konfigürasyon sonrası ilk "
    "birkaç bozuk saat kenarı için bootloader'a bir 'warm-up' (dummy) okuması eklendi.",
    "2) CV32E40P'nin latch-tabanlı kayıt dosyası FPGA'da kötü sentezlenir. Çözüm: build betiğinde "
    "latch sürümü atlanıp flip-flop tabanlı kayıt dosyası kullanıldı.",
    "3) Gerçek S25FL256S flash'ta tCO gecikmesi kaynaklı bit kayması. Çözüm: QSPI RX örnekleme "
    "noktası yükselen kenara taşındı (SPI mod-0 düzeltmesi).",
    "4) SoC'un MMCM kilitlenmeden çalışmaya başlama riski. Çözüm: reset, mmcm_locked sinyaline "
    "bağlandı ve iki kademeli senkronizatörle async-assert/sync-release uygulandı; saat "
    "kararlı olana kadar SoC reset'te tutulur.",
    "5) Build ortamı tuzakları (Vivado'nun WSL/Windows altında çalışması ve çıktı yolları). Çözüm: "
    "build betiği bu ortama uygun olacak şekilde düzenlendi; küresel include başlıkları (axi, "
    "common_cells) Vivado'nun ayrı derleme birimi yaklaşımına göre doğru sırayla eklendi.",
]:
    place(make_para(t, sa=4))

# --- sil + kaydet ---
placeholder._element.getparent().remove(placeholder._element)
doc.save(SRC)
print("OK - 2.4 yazıldı, placeholder silindi.")
