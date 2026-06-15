#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# ============================================
# Ostim BLogic Mikroelektronik
# write_dtr_3.py  -  3. Çip Tasarım Akışı (kısa, plan odaklı)
# ============================================
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'BLogic_MCU_DTR_taslak_01.docx'
doc = docx.Document(SRC)

def _font(run, size, bold=False, italic=False):
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold; run.italic = italic

def make_para(text, size=11, bold=False, italic=False, sb=0, sa=6, justify=True):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(text), size, bold, italic)
    return p

def make_heading(text):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(8); pf.space_after = Pt(2); pf.keep_with_next = True
    _font(p.add_run(text), 12, bold=True)
    return p

def set_cell(cell, text, bold=False, size=10, shade=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    _font(p.add_run(text), size, bold)
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), shade); tcPr.append(shd)

def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        b.append(e)
    tblPr.append(b)

def _width(table, pct=5000):
    tblPr = table._tbl.tblPr
    for ex in tblPr.findall(qn('w:tblW')): tblPr.remove(ex)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), str(pct)); tblPr.append(w)

def make_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _borders(t); _width(t); t.autofit = True
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, bold=True, size=10, shade='D9D9D9')
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(t.rows[i].cells[j], val, size=10)
    return t

heading_p = None; placeholder = None
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') and 'Çip Tasarım Akışı' in p.text:
        heading_p = p
    if p.text.strip().startswith('{Ticari EDA araçlarının'):
        placeholder = p
assert heading_p is not None and placeholder is not None, 'bulunamadı'

anchor = heading_p._element
def place(obj):
    global anchor
    anchor.addnext(obj._element); anchor = obj._element; return obj

# ===================== İÇERİK (kısa) =====================
place(make_para(
    "Ticari EDA araçları DTR aşamasının ardından sağlanacağı için bu aşamada fiziksel tasarım "
    "(ASIC) akışına henüz başlanmamıştır; bu başlık altında planlanan akış aşamaları ve hâlihazırda "
    "tamamlanan öncül çalışmalar özetlenmektedir. Hedeflenen sayısal ASIC akışı aşağıdaki gibidir:"))

place(make_table(
    ["Aşama", "Plan / açıklama"],
    [
     ["1. RTL dondurma + Mantık sentezi",
      "RTL → kapı-seviyesi netlist; PDK standart hücre kütüphanesi ve SDC zamanlama kısıtlarıyla"],
     ["2. DFT / Tarama zinciri",
      "Test edilebilirlik için scan chain ekleme (çekirdekte scan_cg_en girişi hazır)"],
     ["3. Floorplanning",
      "Çekirdek/IO yerleşimi, güç dağıtım ağı (power grid) ve SRAM makro yerleşimi"],
     ["4. Yerleştirme + CTS",
      "Hücre yerleştirme ve saat ağacı sentezi (clock tree synthesis)"],
     ["5. Yönlendirme (Routing)",
      "Sinyal ve güç hatlarının yönlendirilmesi"],
     ["6. İmza (Signoff)",
      "STA (zamanlama), güç analizi, parazitik çıkarım (RC), DRC/LVS denetimleri"],
     ["7. GDSII + tape-out",
      "Fiziksel düzen (layout) çıktısının üretilmesi ve imalat teslimi"],
    ]))

place(make_heading("Öncül Çalışmalar ve Planlar"))
place(make_para(
    "Ticari araçlar henüz sağlanmadığından detaylı fiziksel tasarım faaliyeti yürütülmemiş; ancak "
    "ASIC akışına sağlam bir zemin hazırlanmıştır. Tüm RTL sentezlenebilir biçimde yazılmış ve "
    "Genesys-2 FPGA'sında doğrulanmıştır (bkz. 2.4); bu, mantık sentezi için güçlü bir öncüldür. "
    "Tasarım DFT'ye hazırdır (scan_cg_en girişi mevcuttur) ve yalnız-doğrulamaya yönelik SVA "
    "monitörleri sentezden otomatik hariç tutulur. Davranışsal SRAM modelleri ASIC akışında foundry "
    "SRAM derleyici makrolarıyla değiştirilecek; kayıt dosyası için alan/güç açısından latch-tabanlı "
    "sürüm değerlendirilecektir. Hedef sistem saati 50 MHz'dir. Araçlar sağlandığında sentez → "
    "yerleştirme-bağlama → imza → GDSII adımları yukarıdaki sırayla yürütülecektir."))

place(make_heading("Yararlanılan Topluluklar ve Kanallar"))
place(make_para(
    "Süreç boyunca açık kaynak ve yarışma topluluklarından yararlanılmıştır: çekirdek için OpenHW "
    "Group (CV32E40P) dokümantasyonu ve topluluğu; veri yolu için PULP Platform'un açık AXI ve "
    "common_cells kütüphaneleri; doğrulama için Verilator ve riscv-arch-test toplulukları. Ayrıca "
    "DDK/TEKNOFEST Slack kanalları, demo-testbench ve tasarım akışıyla ilgili sorularda başvurulan "
    "başlıca kaynaklar olmuştur."))

placeholder._element.getparent().remove(placeholder._element)
doc.save(SRC)
print("OK - 3. bölüm yazıldı.")
