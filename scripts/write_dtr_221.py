#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-FileCopyrightText: 2026 BLogic Mikroelektronik - Berk Muammer Kuzu and Berkin Demircan
# SPDX-License-Identifier: GPL-3.0-only
# Licensed under the GNU General Public License version 3 only.
# See the LICENSE file in the repository root for the full license text.

# ============================================
# Ostim BLogic Mikroelektronik
# write_dtr_221.py  -  DTR 2.2.1 bolumunu docx'e yazar
# ============================================
"""2.2.1 İşlemci ve Bus Yapısının Tasarımı (10p) -> BLogic_MCU_DTR_taslak_01.docx"""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'BLogic_MCU_DTR_taslak_01.docx'
doc = docx.Document(SRC)

def _font(run, size, bold=False, italic=False):
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold; run.italic = italic

def make_para(text, size=11, bold=False, italic=False, sb=0, sa=6, justify=True):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(text), size, bold, italic)
    return p

def make_heading(text):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format; pf.space_before = Pt(10); pf.space_after = Pt(3); pf.keep_with_next = True
    _font(p.add_run(text), 12, bold=True)
    return p

def set_cell(cell, text, bold=False, size=10, shade=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    _font(p.add_run(text), size, bold)
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), shade)
        tcPr.append(shd)

def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '808080')
        b.append(e)
    tblPr.append(b)

def _width(table, pct=5000):
    tblPr = table._tbl.tblPr
    for ex in tblPr.findall(qn('w:tblW')): tblPr.remove(ex)
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), str(pct)); tblPr.append(w)

def make_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _borders(t); _width(t); t.autofit = True
    for j, h in enumerate(headers):
        set_cell(t.rows[0].cells[j], h, bold=True, size=10, shade='D9D9D9')
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(t.rows[i].cells[j], val, size=10)
    return t

# heading ve placeholder'i bul
heading_p = None; placeholder = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 3' and 'İşlemci ve Bus Yapısının' in p.text:
        heading_p = p
    if p.text.strip().startswith('{CV32E40P işlemcisinin portları'):
        placeholder = p
assert heading_p is not None and placeholder is not None, 'Heading/placeholder bulunamadı'

anchor = heading_p._element
def place(obj):
    global anchor
    anchor.addnext(obj._element); anchor = obj._element; return obj

# icerik
place(make_para(
    "BLogic MCU'nun işlemci ve veri yolu mimarisi, CV32E40P çekirdeğinin yerel OBI (Open Bus "
    "Interface) arayüzünden başlayıp, performans gereksinimine göre ayrıştırılmış iki katmanlı "
    "bir AXI omurgasına uzanan bir hiyerarşi üzerine kurulmuştur: yüksek hacimli bellek "
    "erişimleri AXI4 üzerinden, düşük hacimli çevre birimi yazmaç erişimleri ise AXI4-Lite "
    "üzerinden taşınmaktadır. Tüm AXI arayüzleri 32-bit adres, 32-bit veri, 4-bit ID ve 1-bit "
    "user genişliğindedir. Bu alt bölümde çekirdek portlarının bağlanışı, ara köprüler, adres "
    "kod-çözme mantığı, seçilen yöntemlerin gerekçeleri ve karşılaşılan zorluklar ile çözümleri "
    "ayrıntılandırılmaktadır. Bu hedef bu aşamada başarılmış olup, işlemci-veri yolu bütünü "
    "hem simülasyonda hem de DDK demo-testbench'inde doğrulanmıştır."))

# cekirdek
place(make_heading("CV32E40P Çekirdeğinin Yapılandırması ve Port Bağlantıları"))
place(make_para(
    "Çekirdek, soc_top içinde cv32e40p_top sarmalayıcısı üzerinden, COREV_PULP=0 ve FPU=0 "
    "parametreleriyle örneklenmiştir; yani PULP özel komut uzantıları ve kayan-nokta birimi "
    "devre dışıdır. Sonuçta saf bir RV32IMC (temel tamsayı I, çarpma/bölme M, sıkıştırılmış C) "
    "makine-modu işlemci elde edilir; A (atomik), PMP ve kullanıcı/süpervizör modları yoktur. "
    "Boru hattı dört aşamalıdır (IF / ID / EX / WB)."))
place(make_para(
    "Çekirdeğin en belirleyici mimari özelliği, biri yalnızca komut getirmeye, diğeri load/store "
    "veri erişimine ayrılmış iki bağımsız OBI portudur. Bu sayede işlemci aynı çevrimde hem "
    "komut çekip hem veri erişimi yapabilir. Bu iki port, ayrı birer OBI→AXI4 köprüsüne "
    "bağlanmaktadır. Aşağıdaki tablo çekirdeğin başlıca portlarını ve bağlandıkları yerleri "
    "özetlemektedir.", sa=4))
place(make_table(
    ["Port / sinyal grubu", "Sinyaller", "Bağlantı / açıklama"],
    [
     ["Buyruk OBI (salt-okunur)",
      "instr_req_o, instr_gnt_i, instr_rvalid_i, instr_addr_o[31:0], instr_rdata_i[31:0]",
      "OBI→AXI4 köprüsü (ID 0); yalnız komut getirme, yazma yok"],
     ["Veri OBI (R/W)",
      "data_req_o, data_gnt_i, data_rvalid_i, data_we_o, data_be_o[3:0], data_addr_o[31:0], data_wdata_o[31:0], data_rdata_i[31:0]",
      "OBI→AXI4 köprüsü (ID 1); load/store"],
     ["Kesme arayüzü",
      "irq_i[31:0], irq_ack_o, irq_id_o[4:0]",
      "irq_i = irq_vector (timer=bit 16, YZ=bit 17, UART_1 akış=bit 18); ack/id çekirdek içinde işlenir"],
     ["Boot / CSR adresleri",
      "boot_addr_i, mtvec_addr_i, dm_halt_addr_i, hart_id_i",
      "boot=0x0000_0000 (Boot ROM), mtvec=0x0001_0000 (Instr SRAM), hart_id=0"],
     ["Hata ayıklama",
      "debug_req_i, debug_*_o",
      "debug_req_i = 1'b0 (JTAG opsiyonel, bu sürümde gerçeklenmedi)"],
     ["Diğer kontrol",
      "fetch_enable_i, pulp_clock_en_i, scan_cg_en_i",
      "fetch_enable=1, clock_en=1, scan_cg_en=0 (tarama yalnız ASIC DFT'de)"],
    ]))

# kopru ve interconnect
place(make_heading("OBI → AXI4 Köprüleri ve Çekirdeğin Veri Yoluna Bağlanışı"))
place(make_para(
    "CV32E40P'nin yerel OBI arayüzü tek fazlı bir istek/onay (req/gnt) artı ayrı bir yanıt "
    "(rvalid) protokolüyken, sistem veri yolu beş bağımsız kanallı AXI'dir. Bu uyumsuzluğu "
    "gidermek için özgün bir obi_to_axi köprüsü tasarlanmıştır. Köprü iki kez örneklenir: "
    "buyruk portu için ID 0 (yazma tarafı kapalı: we=0, be=4'b1111, wdata=0), veri portu için "
    "ID 1 (tam okuma/yazma). Köprü, altı durumlu bir durum makinesiyle (IDLE, WAIT_AW_W, "
    "WAIT_AW, WAIT_W, WAIT_B, WAIT_R) çalışır: yazmada AW ve W kanallarını eşzamanlı sürer ve "
    "her kanalın kabulünü ayrı bayraklarla takip eder; AXI yanıtını (B veya R) çekirdeğin OBI "
    "rvalid'ine geri taşır. Her işlem tek-beat'tir (aw_len/ar_len=0, size=4 bayt, INCR) ve "
    "köprü tek-bekleyen (single-outstanding) tasarlanmıştır; yani bir işlem tamamlanmadan "
    "yenisi başlatılmaz."))
place(make_heading("AXI4 Çapraz-Bağlama (soc_axi_interconnect) ve Adres Kod-Çözme"))
place(make_para(
    "İki AXI4 master (buyruk, veri), elle yazılmış özgün bir AXI4 çapraz-bağlama olan "
    "soc_axi_interconnect içinde beş slave'e yönlendirilir: Boot ROM, Buyruk SRAM, Veri SRAM, "
    "AI SRAM (arbiter üzerinden) ve çevre birimi portu. Yönlendirme, döner-sıralı (round-robin) "
    "bir hakem yerine saf kombinasyonel adres kod-çözme ile yapılır; buyruk ve veri yolları "
    "statik olarak ayrı slave kanallarına bağlandığı için master'lar arasında çekişme oluşmaz "
    "ve ek bir hakem birimine gerek kalmaz. Buyruk yolu yalnızca okuma yapar. Kod-çözme mantığı "
    "şu şekildedir:", sa=4))
place(make_table(
    ["Yol", "Koşul (adres)", "Hedef"],
    [
     ["Buyruk (okuma)", "addr[19:16] == 0x0", "Boot ROM (0x0000_xxxx)"],
     ["Buyruk (okuma)", "addr[19:16] != 0x0", "Instruction SRAM (0x0001_xxxx)"],
     ["Veri (yazma)", "addr[31:28] == 0x4", "Çevre birimleri (0x4000_xxxx)"],
     ["Veri (yazma)", "addr[19:16] == 0x1", "Instruction SRAM (bootloader yazması)"],
     ["Veri (yazma)", "addr[19:16] == 0x3", "AI SRAM (arbiter)"],
     ["Veri (yazma)", "diğer", "Data SRAM (0x0002_xxxx)"],
     ["Veri (okuma)", "0x4 / [19:16]=0x3 / diğer", "Çevre / AI SRAM / Data SRAM (Buyruk SRAM veri yolundan OKUNAMAZ)"],
    ]))
place(make_heading("Çevre Birimlerinin Bağlanışı: AXI4→AXI-Lite Köprüsü ve periph_decoder"))
place(make_para(
    "Çevre portundaki AXI4-full akış, axi4_to_axilite_bridge ile AXI4-Lite'a indirgenir. Bu "
    "köprü durum makinesi içermeyen, gecikmesiz (kombinasyonel) bir geçiştir: AXI4'e özgü "
    "sinyalleri (len, size, burst, lock, cache, prot, qos) düşürür, yanıt tarafında ID'leri "
    "yeniden iliştirir (s_bid=s_awid, s_rid=s_arid) ve her erişimi tek-beat olarak ele alır. "
    "Ardından periph_decoder, adresin [11:8] alanına bakarak 7 çevre birimini (UART_0, GPIO, "
    "Timer, UART_1, I2C, QSPI, YZ-CSR) 0x100 baytlık pencerelerle seçer. Tüm çevre birimleri "
    "aynı AXI4-Lite slave şablonunu kullanır; böylece register erişimleri tek-tip ve "
    "öngörülebilir olur. Haritalanmamış bir adrese erişimde çözücü, yanıt kodu olarak DECERR "
    "(2'b11) ve okuma verisinde 0xDEADBEEF üreterek erişim hatasını yazılıma bildirir."))

# arti/eksi
place(make_heading("Seçilen Yöntemlerin Artı ve Eksileri"))
place(make_table(
    ["Tasarım kararı", "Artıları", "Eksileri / risk"],
    [
     ["AXI4 + AXI4-Lite (APB/Wishbone yerine)",
      "Endüstri standardı, şartname EK uyumlu, ayrık kanallar, geniş IP ve SVA ekosistemi",
      "APB'ye göre daha karmaşık; iki katmanlı yapı ile bu maliyet düşük tutuldu"],
     ["İki katmanlı veri yolu (tek yol yerine)",
      "Bellek tarafı yüksek başarımlı (AXI4), çevre tarafı sade (AXI-Lite); register erişimi için yeterli",
      "Ek köprü bloğu; ancak kombinasyonel olduğundan gecikmesi ~0"],
     ["Özel elle-yazılmış interconnect (hazır xbar IP yerine)",
      "Minimal alan, tam kontrol, kanal ayrımıyla hakemsiz/deterministik yönlendirme",
      "Ölçeklenince elle bakım zorlaşır; 2 master × 5 slave sabit topolojide kabul edilebilir"],
     ["Tek-bekleyen OBI→AXI köprü",
      "Basit, küçük alan, doğrulaması kolay",
      "Aynı anda tek işlem → başarım sınırı; çekirdek çoğunlukla tek-beat eriştiği için pratik etki düşük, ileride artırılabilir"],
    ]))

# zorluklar
place(make_heading("Karşılaşılan Zorluklar ve Çözümleri"))
for t in [
    "1) OBI ↔ AXI protokol uyumsuzluğu: OBI tek fazlı req/gnt + ayrı rvalid kullanırken AXI beş "
    "kanallıdır. Çözüm: AW ve W kanallarını eşzamanlı süren, her kanalın kabulünü ayrı "
    "bayraklarla izleyen ve B/R yanıtını OBI rvalid'e geri taşıyan özel bir FSM köprü tasarlandı.",
    "2) Buyruk SRAM'ine hem yazma (bootloader) hem getirme (CPU) ihtiyacı, tek-portlu davranışsal "
    "bir RAM ile karşılanmak zorundaydı. Çözüm: çift-portlu bir donanım makrosu kullanmak yerine, "
    "interconnect seviyesinde kanallar ayrıştırıldı; getirme yalnız buyruk yolundan (okuma), "
    "yazma yalnız veri yolundan yapılır. Böylece sözde çift-port davranışı elde edildi.",
    "3) Sabit verilerin (.rodata) erişilebilirliği: CPU veri portunun okuma yolu Buyruk SRAM'e "
    "açık değildir. Çözüm: linker betiğinde .rodata bölümü Veri SRAM'e yerleştirildi; bu, "
    "donanım-yazılım birlikte tasarımının somut bir örneğidir.",
    "4) Hatalı/tanımsız adres erişimlerinin görünürlüğü: Çözüm: periph_decoder, eşleşmeyen "
    "adreslerde DECERR (2'b11) ve 0xDEADBEEF döndürerek hatayı yazılıma bildirir; bu, hata "
    "ayıklamayı belirgin biçimde kolaylaştırır.",
    "5) AI SRAM'e çoklu master erişimi (CPU, YZ hızlandırıcı, UART_1 DMA): Çözüm: ana veri yolunu "
    "basit tutmak için erişim, ayrı bir ai_sram_arbiter ile çözüldü (öncelik: YZ > akış > CPU). "
    "Ayrıntısı 2.2.3'te verilmiştir.",
    "6) Master ayrımı ve yanıt yönlendirme: buyruk ve veri master'larına farklı AXI ID (0 ve 1) "
    "atanarak yanıtların doğru master'a dönmesi garanti altına alındı.",
]:
    place(make_para(t, sa=4))

# dogrulama
place(make_heading("Doğrulama ve Hedefin Başarılması"))
place(make_para(
    "İşlemci-veri yolu bütünü bu aşamada hedeflendiği gibi çalışır durumdadır ve kapsamlı bir "
    "doğrulama akışından geçmiştir: Verilator tabanlı regresyon takımı boot akışını, SoC-YZ veri "
    "yolunu, UART-stream ve QSPI senaryolarını kendi kendini denetleyen testlerle koşturmakta; "
    "çekirdek davranışı Spike komut kümesi simülatörüyle lockstep komut izi karşılaştırmasına tabi "
    "tutulmakta; tüm AXI-Lite çevre arayüzleri ile AXI4 master'ları üzerinde SVA protokol "
    "monitörleri sürekli denetim yapmakta ve riscv-arch-test uyumluluk paketi çalıştırılmaktadır. "
    "Ayrıca yarışma kapsamında istenen DDK demo-testbench senaryosu (UART 'hello world') Vivado "
    "2021.2 ortamında başarıyla geçmiş ve tasarım Genesys-2 kartında bring-up edilmiştir. Bu "
    "sonuçlar, işlemci ve veri yolu tasarımı hedefinin bu aşamada başarıldığını göstermektedir."))

# placeholder'i sil ve kaydet
placeholder._element.getparent().remove(placeholder._element)
doc.save(SRC)
print("OK - 2.2.1 yazıldı, placeholder silindi.")
